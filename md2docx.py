#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""把 技术文档.md 转成 技术文档.docx。

可重跑：把 figs/fig1..fig6、fig9 换成真实绘制图后，重新运行本脚本即可刷新嵌入。
"""
import os
import re

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Inches, Pt, RGBColor

BASE = os.path.dirname(os.path.abspath(__file__))
MD = os.path.join(BASE, "技术文档.md")
DOCX = os.path.join(BASE, "技术文档.docx")

EA = "宋体"                # 正文：宋体
EA_HEAD = "黑体"           # 标题：黑体
LATIN = "Times New Roman"  # 西文/数字
MONO = "Consolas"
HEADING_SIZE = {0: 18, 1: 16, 2: 14}   # 小二 / 三号 / 四号


def set_run_ea(run, latin=None, ea=EA):
    rpr = run._element.get_or_add_rPr()
    rFonts = rpr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rpr.append(rFonts)
    if latin:
        run.font.name = latin
        rFonts.set(qn("w:ascii"), latin)
        rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), ea)


def shade_paragraph(p, fill="F2F2F2"):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def box_paragraph(p, color="999999"):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + edge)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:space"), "4")
        e.set(qn("w:color"), color)
        pbdr.append(e)
    pPr.append(pbdr)


def latex_to_text(s):
    s = s.replace("\\text{", "").replace("\\lVert", "‖").replace("\\rVert", "‖")
    s = s.replace("\\min", "min").replace("\\in", "∈").replace("\\_", "_")
    s = s.replace("\\ ", " ").replace("{", "").replace("}", "")
    return s


INLINE = re.compile(r"(\*\*.*?\*\*|`[^`]*`|\$[^$]*\$|\[[^\]]*\]\([^)]*\))")


def add_inline(p, text):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            r = p.add_run(tok[2:-2])
            r.bold = True
        elif tok.startswith("`") and tok.endswith("`") and len(tok) > 2:
            r = p.add_run(tok[1:-1])
            r.font.name = MONO
            r.font.size = Pt(9.5)
            set_run_ea(r, latin=MONO)
        elif tok.startswith("$") and tok.endswith("$") and len(tok) > 2:
            r = p.add_run(latex_to_text(tok[1:-1]))
            r.italic = True
        elif tok.startswith("[") and ")](" not in tok and "](" in tok:
            txt = tok[1:tok.index("](")]
            r = p.add_run(txt)
            r.font.color.rgb = RGBColor(0x05, 0x63, 0xC1)
        else:
            p.add_run(tok)
    for r in p.runs:
        set_run_ea(r)


doc = Document()

# 基础样式：正文宋体小四 12pt、西文 Times New Roman、1.5 倍行距
normal = doc.styles["Normal"]
normal.font.name = LATIN
normal.font.size = Pt(12)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.5
rpr = normal.element.get_or_add_rPr()
rF = rpr.find(qn("w:rFonts"))
if rF is None:
    rF = OxmlElement("w:rFonts")
    rpr.append(rF)
rF.set(qn("w:ascii"), LATIN)
rF.set(qn("w:hAnsi"), LATIN)
rF.set(qn("w:eastAsia"), EA)

# 页面边距：上下 2.5cm、左右 3.0cm
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(3.0)


def add_heading(text, level):
    if level == 0:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p = doc.add_paragraph(style="Heading %d" % level)
    add_inline(p, text)
    for r in p.runs:
        r.bold = True
        r.font.size = Pt(HEADING_SIZE[level])
        r.font.color.rgb = RGBColor(0, 0, 0)
        set_run_ea(r, latin=LATIN, ea=EA_HEAD)
    p.paragraph_format.space_before = Pt(12 if level == 0 else 6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.3
    return p


def first_line_indent(p, chars=2):
    pPr = p._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars * 100))


def add_para(text):
    p = doc.add_paragraph()
    if text.startswith("*") and text.endswith("*") and not text.startswith("**"):
        text = text[1:-1]
        p.add_run(text).italic = True
        for r in p.runs:
            set_run_ea(r)
    else:
        add_inline(p, text)
    first_line_indent(p)
    return p


def add_blockquote(text):
    p = doc.add_paragraph()
    add_inline(p, text)
    if text.startswith("**图"):
        # 图注：宋体五号 10.5pt、居中、无底色
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.0
        for r in p.runs:
            r.font.size = Pt(10.5)
            set_run_ea(r, latin=LATIN, ea=EA)
    else:
        p.paragraph_format.left_indent = Inches(0.25)
        shade_paragraph(p, "F5F5F5")
    return p


def add_hr():
    p = doc.add_paragraph()
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pbdr.append(bottom)
    pPr.append(pbdr)


def add_code_block(code, lang):
    label = lang if lang else "code"
    p = doc.add_paragraph()
    r = p.add_run("[%s]" % label)
    r.font.name = MONO
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
    set_run_ea(r, latin=MONO)
    p.paragraph_format.space_after = Pt(0)
    for ln in code:
        p = doc.add_paragraph()
        r = p.add_run(ln if ln else " ")
        r.font.name = MONO
        r.font.size = Pt(9)
        set_run_ea(r, latin=MONO)
        p.paragraph_format.space_after = Pt(0)
        shade_paragraph(p, "F2F2F2")


def split_row(ln):
    ln = ln.strip()
    if ln.startswith("|"):
        ln = ln[1:]
    if ln.endswith("|"):
        ln = ln[:-1]
    return [c.replace("\\|", "|").strip() for c in re.split(r"(?<!\\)\|", ln)]


def parse_table(block):
    header = split_row(block[0])
    rows = [split_row(ln) for ln in block[2:]]
    return header, rows


def add_table(header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    hdr = t.rows[0].cells
    for i, h in enumerate(header):
        hp = hdr[i].paragraphs[0]
        add_inline(hp, h)
        for r in hp.runs:
            r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, c in enumerate(row):
            cp = cells[i].paragraphs[0]
            add_inline(cp, c)
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                para.paragraph_format.line_spacing = 1.0
                for r in para.runs:
                    r.font.size = Pt(10.5)
                    set_run_ea(r, latin=LATIN, ea=EA)


def add_image(path, alt):
    full = os.path.join(BASE, path)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if os.path.exists(full):
        with Image.open(full) as im:
            iw, ih = im.size
        ratio = iw / ih
        max_w, max_h = 5.5, 7.0
        if ratio >= max_w / max_h:
            p.add_run().add_picture(full, width=Inches(max_w))
        else:
            p.add_run().add_picture(full, height=Inches(max_h))
    else:
        r = p.add_run("[ 图片占位 ]  " + alt)
        r.italic = True
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        set_run_ea(r)
        box_paragraph(p)


def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    add_inline(p, text)
    return p


def add_ordered(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    add_inline(p, text)
    return p


lines = open(MD, encoding="utf-8").read().split("\n")
i, n = 0, len(lines)
while i < n:
    s = lines[i].strip()
    if s == "":
        i += 1
        continue
    if s == "---":
        add_hr()
        i += 1
        continue
    if s.startswith("```"):
        lang = s[3:].strip()
        i += 1
        code = []
        while i < n and not lines[i].strip().startswith("```"):
            code.append(lines[i])
            i += 1
        i += 1  # 跳过结束 ```
        add_code_block(code, lang)
        continue
    if s.startswith("### "):
        add_heading(s[4:].strip(), 2)
        i += 1
        continue
    if s.startswith("## "):
        add_heading(s[3:].strip(), 1)
        i += 1
        continue
    if s.startswith("# "):
        add_heading(s[2:].strip(), 0)
        i += 1
        continue
    if s.startswith("|"):
        block = []
        while i < n and lines[i].strip().startswith("|"):
            block.append(lines[i].strip())
            i += 1
        header, rows = parse_table(block)
        add_table(header, rows)
        continue
    if s.startswith(">"):
        add_blockquote(s[1:].strip())
        i += 1
        continue
    if s.startswith("!["):
        m = re.match(r"!\[([^\]]*)\]\(([^)]*)\)", s)
        add_image(m.group(2), m.group(1))
        i += 1
        continue
    if re.match(r"^\d+\.\s", s):
        add_ordered(s)
        i += 1
        continue
    if s.startswith("- "):
        add_bullet(s[2:])
        i += 1
        continue
    add_para(s)
    i += 1

doc.save(DOCX)
print("saved:", DOCX)
